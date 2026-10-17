"""Normalisasi tabel dan ekspor ke Markdown / CSV / XLSX.

Semua engine menghasilkan tabel dalam bentuk berbeda:
  * pdfplumber  -> list[list[str|None]] dengan sel bergabung bernilai None
  * PP-StructureV3 -> HTML (<table>, dengan colspan/rowspan)
  * Docling     -> TableItem (punya export_to_dataframe/markdown)
  * Marker      -> HTML di dalam blok

Modul ini menyatukan semuanya ke satu representasi: grid persegi list[list[str]].
Sel yang digabung (merged) di-"isi-turun"/"isi-kanan" supaya setiap sel grid
punya nilai — ini penting untuk RAG: tabel kalender akademik dengan sel
"Semester Ganjil" yang membentang 5 baris akan kehilangan konteks kalau
sel-sel turunannya dibiarkan kosong.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Iterable, Sequence

from .models import TableData
from .utils import collapse_to_single_line, get_logger, normalize_cell

_NUMERIC_RE = re.compile(r"^[\s\d.,%()+\-/:]+$")


# --------------------------------------------------------------------------- #
# Normalisasi grid
# --------------------------------------------------------------------------- #


def normalize_grid(
    rows: Sequence[Sequence[Any]],
    *,
    fill_merged: bool = True,
    drop_empty_rows: bool = True,
    drop_empty_cols: bool = True,
) -> tuple[list[list[str]], list[str]]:
    """Ubah grid mentah menjadi grid persegi berisi string bersih.

    Returns:
        (grid, daftar_peringatan)
    """
    warnings: list[str] = []
    grid: list[list[str]] = [[normalize_cell(c) for c in row] for row in rows]
    grid = [r for r in grid if r is not None]
    if not grid:
        return [], ["tabel kosong"]

    widths = {len(r) for r in grid}
    if len(widths) > 1:
        warnings.append(
            f"jumlah kolom tidak konsisten antar baris ({sorted(widths)}); "
            "baris pendek dilengkapi sel kosong"
        )
    n_cols = max(widths)
    grid = [r + [""] * (n_cols - len(r)) for r in grid]

    if fill_merged:
        _fill_down(grid)

    if drop_empty_rows:
        kept = [r for r in grid if any(c.strip() for c in r)]
        if len(kept) != len(grid):
            warnings.append(f"{len(grid) - len(kept)} baris kosong dibuang")
        grid = kept

    if drop_empty_cols and grid:
        keep = [
            i
            for i in range(len(grid[0]))
            if any(row[i].strip() for row in grid)
        ]
        if len(keep) != len(grid[0]):
            warnings.append(f"{len(grid[0]) - len(keep)} kolom kosong dibuang")
            grid = [[row[i] for i in keep] for row in grid]

    return grid, warnings


def _fill_down(grid: list[list[str]]) -> None:
    """Isi sel kosong dengan nilai di atasnya HANYA bila pola merge jelas.

    Heuristik konservatif: sebuah kolom hanya di-isi-turun jika kolom itu
    memiliki setidaknya satu nilai dan >= 30% selnya kosong — ciri khas sel
    yang digabung secara vertikal. Tanpa pagar ini, kolom yang memang
    sebagian kosong (mis. kolom "Keterangan") akan terisi nilai palsu.
    """
    if not grid:
        return
    n_rows, n_cols = len(grid), len(grid[0])
    if n_rows < 3:
        return
    for col in range(n_cols):
        values = [grid[r][col] for r in range(n_rows)]
        n_empty = sum(1 for v in values if not v.strip())
        n_filled = n_rows - n_empty
        if n_filled == 0 or n_empty / n_rows < 0.30:
            continue
        # Jangan isi-turun kolom numerik (nilai angka tidak diwarisi).
        filled_vals = [v for v in values if v.strip()]
        if filled_vals and all(_NUMERIC_RE.match(v) for v in filled_vals):
            continue
        last = ""
        for r in range(n_rows):
            if grid[r][col].strip():
                last = grid[r][col]
            elif last:
                grid[r][col] = last


def guess_header_rows(grid: Sequence[Sequence[str]]) -> int:
    """Tebak jumlah baris header (0, 1, atau 2).

    Sinyal: baris header biasanya seluruhnya non-numerik sementara badan
    tabel memuat angka/tanggal. Dokumen akademik juga sering memakai header
    dua tingkat ("Semester | Ganjil | Genap").
    """
    if not grid:
        return 0
    if len(grid) == 1:
        return 1

    def non_numeric_ratio(row: Sequence[str]) -> float:
        cells = [c for c in row if c.strip()]
        if not cells:
            return 0.0
        return sum(1 for c in cells if not _NUMERIC_RE.match(c)) / len(cells)

    r0 = non_numeric_ratio(grid[0])
    if r0 < 0.6:
        return 0  # baris pertama sudah berisi data

    if len(grid) >= 4:
        r1 = non_numeric_ratio(grid[1])
        body = [non_numeric_ratio(r) for r in grid[2:6]]
        avg_body = sum(body) / len(body) if body else 1.0
        # Baris ke-2 masih "seperti header" sementara badan tabel jelas berbeda.
        if r1 >= 0.85 and avg_body < 0.7:
            return 2
    return 1


def build_table(
    rows: Sequence[Sequence[Any]],
    *,
    page: int,
    index: int,
    engine: str,
    caption: str | None = None,
    bbox: tuple[float, float, float, float] | None = None,
    n_header_rows: int | None = None,
) -> TableData:
    grid, warns = normalize_grid(rows)
    n_head = guess_header_rows(grid) if n_header_rows is None else n_header_rows
    return TableData(
        rows=grid,
        n_header_rows=max(0, min(n_head, len(grid))),
        page=page,
        index=index,
        caption=caption,
        engine=engine,
        bbox=bbox,
        warnings=warns,
    )


# --------------------------------------------------------------------------- #
# Konversi dari HTML (PP-StructureV3, Marker, Docling-HTML)
# --------------------------------------------------------------------------- #


def html_table_to_rows(html: str) -> list[list[str]]:
    """Parse satu <table> HTML menjadi grid persegi, menghormati colspan/rowspan.

    Memakai lxml bila tersedia (toleran terhadap HTML rusak), jatuh ke
    html.parser bawaan Python bila tidak.
    """
    try:
        from lxml import html as lhtml  # type: ignore

        root = lhtml.fromstring(html)
        tr_nodes = root.xpath("//tr")
        raw: list[list[tuple[str, int, int]]] = []
        for tr in tr_nodes:
            cells: list[tuple[str, int, int]] = []
            for cell in tr.xpath("./td|./th"):
                text = cell.text_content()
                cs = _to_int(cell.get("colspan"), 1)
                rs = _to_int(cell.get("rowspan"), 1)
                cells.append((text, cs, rs))
            raw.append(cells)
    except ImportError:  # pragma: no cover
        raw = _parse_table_stdlib(html)

    return _expand_spans(raw)


def _to_int(value: Any, default: int = 1) -> int:
    try:
        n = int(str(value).strip())
        return n if 1 <= n <= 100 else default
    except (TypeError, ValueError):
        return default


def _expand_spans(raw: list[list[tuple[str, int, int]]]) -> list[list[str]]:
    """Bentangkan colspan/rowspan menjadi grid penuh (nilai disalin)."""
    grid: dict[tuple[int, int], str] = {}
    occupied: set[tuple[int, int]] = set()
    max_col = 0
    for r, row in enumerate(raw):
        c = 0
        for text, cs, rs in row:
            while (r, c) in occupied:
                c += 1
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
                    grid[(r + dr, c + dc)] = text
            c += cs
            max_col = max(max_col, c)
    # Jumlah baris = indeks baris tertinggi yang benar-benar terisi + 1.
    # (rowspan sudah memperluas occupied, jadi cukup lihat kunci grid.)
    n_rows = max((r for r, _ in grid), default=-1) + 1
    return [
        [normalize_cell(grid.get((r, c), "")) for c in range(max_col)]
        for r in range(n_rows)
    ]


def _parse_table_stdlib(html: str) -> list[list[tuple[str, int, int]]]:  # pragma: no cover
    from html.parser import HTMLParser

    class _P(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.rows: list[list[tuple[str, int, int]]] = []
            self._row: list[tuple[str, int, int]] | None = None
            self._buf: list[str] = []
            self._attrs: dict[str, Any] = {}
            self._in_cell = False

        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            if tag == "tr":
                self._row = []
            elif tag in ("td", "th"):
                self._in_cell = True
                self._buf = []
                self._attrs = dict(attrs)
            elif tag == "br" and self._in_cell:
                self._buf.append("\n")

        def handle_endtag(self, tag: str) -> None:
            if tag in ("td", "th") and self._row is not None:
                self._row.append(
                    (
                        "".join(self._buf),
                        _to_int(self._attrs.get("colspan"), 1),
                        _to_int(self._attrs.get("rowspan"), 1),
                    )
                )
                self._in_cell = False
            elif tag == "tr" and self._row is not None:
                self.rows.append(self._row)
                self._row = None

        def handle_data(self, data: str) -> None:
            if self._in_cell:
                self._buf.append(data)

    p = _P()
    p.feed(html)
    return p.rows


def dataframe_to_rows(df: Any) -> list[list[str]]:
    """Ubah pandas DataFrame (dari Docling) menjadi grid termasuk baris header."""
    header = [normalize_cell(c) for c in df.columns.tolist()]
    body = [[normalize_cell(v) for v in row] for row in df.astype(str).values.tolist()]
    # DataFrame dengan kolom generik 0..n bukan header sungguhan.
    if all(h.isdigit() or not h for h in header):
        return body
    return [header, *body]


# --------------------------------------------------------------------------- #
# Ekspor
# --------------------------------------------------------------------------- #


def table_to_markdown(table: TableData, *, max_cell_chars: int = 400) -> str:
    """Render tabel sebagai pipe-table Markdown (GFM).

    Baris baru di dalam sel diganti `<br>` karena pipe-table tidak boleh
    mengandung newline literal. Pipe di dalam isi sel di-escape.
    """
    if table.is_degenerate():
        return ""
    grid = table.rows
    n_cols = table.n_cols
    n_head = table.n_header_rows

    def fmt(cell: str) -> str:
        s = cell.replace("|", "\\|").replace("\n", "<br>")
        s = collapse_to_single_line(s) if "<br>" not in s else s
        if len(s) > max_cell_chars:
            s = s[: max_cell_chars - 1].rstrip() + "…"
        return s

    lines: list[str] = []
    if n_head > 0:
        header = table.header or [""] * n_cols
        header = (header + [""] * n_cols)[:n_cols]
        # Markdown mewajibkan header tidak seluruhnya kosong agar tabel dirender.
        if not any(h.strip() for h in header):
            header = [f"Kolom {i + 1}" for i in range(n_cols)]
        lines.append("| " + " | ".join(fmt(h) for h in header) + " |")
        lines.append("| " + " | ".join(["---"] * n_cols) + " |")
        body = table.body
    else:
        lines.append("| " + " | ".join(f"Kolom {i + 1}" for i in range(n_cols)) + " |")
        lines.append("| " + " | ".join(["---"] * n_cols) + " |")
        body = grid

    for row in body:
        row = (list(row) + [""] * n_cols)[:n_cols]
        lines.append("| " + " | ".join(fmt(c) for c in row) + " |")

    out = "\n".join(lines)
    if table.caption:
        out = f"**{table.caption.strip()}**\n\n{out}"
    return out


def table_to_text_lines(table: TableData) -> str:
    """Representasi "kolom: nilai" per baris.

    Berguna sebagai teks pendamping untuk embedding: dense retriever sering
    lebih mudah mencocokkan "Tanggal: 1 September 2026" daripada sel tabel
    yang berdiri sendiri.
    """
    header = table.header
    lines: list[str] = []
    for row in table.body:
        parts = []
        for i, cell in enumerate(row):
            cell = collapse_to_single_line(cell)
            if not cell:
                continue
            label = collapse_to_single_line(header[i]) if i < len(header) else f"kolom{i + 1}"
            parts.append(f"{label}: {cell}" if label else cell)
        if parts:
            lines.append("; ".join(parts))
    return "\n".join(lines)


def shard_table(table: TableData, rows_per_shard: int) -> list[TableData]:
    """Pecah tabel panjang menjadi beberapa bagian, HEADER DIULANG tiap bagian.

    Mengulang header wajib: tanpa itu, pecahan ke-2 dan seterusnya kehilangan
    nama kolom dan menjadi tidak bermakna saat diambil oleh retriever.
    """
    body = table.body
    if rows_per_shard <= 0 or len(body) <= rows_per_shard:
        return [table]
    head_rows = table.rows[: table.n_header_rows]
    shards: list[TableData] = []
    total = (len(body) + rows_per_shard - 1) // rows_per_shard
    for i in range(0, len(body), rows_per_shard):
        part = i // rows_per_shard + 1
        cap = table.caption or "Tabel"
        shards.append(
            TableData(
                rows=[*head_rows, *body[i : i + rows_per_shard]],
                n_header_rows=table.n_header_rows,
                page=table.page,
                index=table.index,
                caption=f"{cap} (bagian {part} dari {total})",
                engine=table.engine,
                bbox=table.bbox,
                warnings=list(table.warnings),
            )
        )
    return shards


def write_table_csv(table: TableData, path: str | Path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig agar Excel di Windows membaca karakter Indonesia dengan benar.
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        writer.writerows(table.rows)
    return p


def write_tables_xlsx(
    tables: Iterable[TableData],
    path: str | Path,
    *,
    doc_title: str = "",
    doc_id: str = "",
) -> Path | None:
    """Tulis semua tabel satu dokumen ke satu XLSX multi-sheet.

    Sheet pertama ("INDEKS") memuat daftar tabel beserta halaman asalnya,
    sehingga berkas ini bisa dipakai langsung sebagai lampiran skripsi.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tables = [t for t in tables if not t.is_degenerate()]
    if not tables:
        return None

    log = get_logger()
    wb = Workbook()
    idx = wb.active
    idx.title = "INDEKS"
    idx.append(["Dokumen", doc_title or doc_id])
    idx.append(["doc_id", doc_id])
    idx.append([])
    idx.append(["Sheet", "Halaman", "Baris", "Kolom", "Keterangan", "Engine", "Catatan"])
    for cell in idx[4]:
        cell.font = Font(bold=True)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    used: set[str] = set()

    for t in tables:
        name = t.slug
        suffix = 1
        while name in used:
            suffix += 1
            name = f"{t.slug}_{suffix}"
        used.add(name)
        ws = wb.create_sheet(title=name[:31])

        for row in t.rows:
            ws.append([c.replace("\n", " ") for c in row])

        for r in range(1, max(1, t.n_header_rows) + 1):
            if r > ws.max_row:
                break
            for cell in ws[r]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        if t.n_header_rows:
            ws.freeze_panes = ws.cell(row=t.n_header_rows + 1, column=1)

        for col in range(1, t.n_cols + 1):
            longest = max(
                (len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)),
                default=10,
            )
            ws.column_dimensions[get_column_letter(col)].width = min(max(10, longest + 2), 60)

        idx.append(
            [
                name,
                t.page,
                t.n_rows,
                t.n_cols,
                (t.caption or "")[:120],
                t.engine,
                "; ".join(t.warnings)[:200],
            ]
        )

    for col, width in zip("ABCDEFG", (16, 10, 8, 8, 46, 14, 46)):
        idx.column_dimensions[col].width = width

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(p))
    log.debug("menulis %d tabel ke %s", len(tables), p.name)
    return p
